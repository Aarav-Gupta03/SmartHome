import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_docx():
    doc = docx.Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Styles
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t = title_p.add_run("Smart Home IoT Data Engine\n")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(24)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(26, 54, 93) # Navy blue
    
    run_sub = title_p.add_run("Project Functionalities & Review 3 Documentation")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(74, 85, 104)
    
    intro_p = doc.add_paragraph()
    intro_p.paragraph_format.space_before = Pt(10)
    intro_p.paragraph_format.space_after = Pt(15)
    run_intro = intro_p.add_run(
        "This document outlines the core capabilities of the Smart Home IoT Data Engine for Review 3, "
        "divided into User-Defined Functionalities (Application Tier) and NoSQL Functionalities (Database Tier). "
        "All functionalities presented here are brand-new and distinct from Review 2."
    )
    run_intro.font.size = Pt(10.5)
    run_intro.font.italic = True
    run_intro.font.color.rgb = RGBColor(74, 85, 104)

    # Read markdown
    with open("Review_3_Project_Functionalities.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    # Split into sections
    parts = re.split(r'###\s+Part\s+(\d+):\s+(.+)', md_text)
    # parts: [preamble, '1', 'User-Defined...', text1, '2', 'NoSQL...', text2]
    
    i = 1
    while i < len(parts):
        part_num = parts[i]
        part_title = parts[i+1].strip()
        part_body = parts[i+2]
        i += 3
        
        # Add Part heading
        part_h = doc.add_paragraph()
        part_h.paragraph_format.space_before = Pt(20)
        part_h.paragraph_format.space_after = Pt(10)
        p_run = part_h.add_run(f"Part {part_num}: {part_title}")
        p_run.font.name = "Arial"
        p_run.font.size = Pt(16)
        p_run.font.bold = True
        p_run.font.color.rgb = RGBColor(43, 108, 176)
        
        # Split into features (#### N. ...)
        feature_blocks = re.split(r'####\s+(\d+\..+)', part_body)
        f_idx = 1
        while f_idx < len(feature_blocks):
            feat_title = feature_blocks[f_idx].strip()
            feat_content = feature_blocks[f_idx+1]
            f_idx += 2
            
            # Feature heading
            fh = doc.add_paragraph()
            fh.paragraph_format.space_before = Pt(14)
            fh.paragraph_format.space_after = Pt(4)
            fh.paragraph_format.keep_with_next = True
            fh_run = fh.add_run(feat_title)
            fh_run.font.name = "Arial"
            fh_run.font.size = Pt(12.5)
            fh_run.font.bold = True
            fh_run.font.color.rgb = RGBColor(45, 55, 72)
            
            # Extract code block if any
            code_match = re.search(r'```(?:java)?\n(.*?)```', feat_content, re.DOTALL)
            code_text = code_match.group(1).strip() if code_match else None
            
            # Extract image if any
            img_match = re.search(r'!\[(.*?)\]\((.*?)\)', feat_content)
            img_caption = img_match.group(1) if img_match else None
            img_path = img_match.group(2) if img_match else None
            
            # Extract pure description text (remove code and image)
            clean_desc = feat_content
            if code_match:
                clean_desc = clean_desc.replace(code_match.group(0), "")
            if img_match:
                clean_desc = clean_desc.replace(img_match.group(0), "")
            clean_desc = clean_desc.strip()
            
            if clean_desc:
                dp = doc.add_paragraph()
                dp.paragraph_format.space_after = Pt(8)
                d_run = dp.add_run(clean_desc)
                d_run.font.name = "Calibri"
                d_run.font.size = Pt(10.5)
                d_run.font.color.rgb = RGBColor(50, 50, 50)
            
            # If code exists, add as styled callout table
            if code_text:
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                cell.width = Inches(6.8)
                
                # Light grey background
                shading_elm = parse_xml(r'<w:shd {} w:fill="F7FAFC"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shading_elm)
                
                # Border
                borders_elm = parse_xml(r'<w:tcBorders {}><w:left w:val="single" w:sz="24" w:space="0" w:color="3182CE"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(borders_elm)
                
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(4)
                cp.paragraph_format.space_after = Pt(4)
                c_run = cp.add_run(code_text)
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(9)
                c_run.font.color.rgb = RGBColor(45, 55, 72)
                
                sp_p = doc.add_paragraph()
                sp_p.paragraph_format.space_after = Pt(4)
                
            # Add image if exists
            if img_path and os.path.exists(img_path):
                ip = doc.add_paragraph()
                ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ip.paragraph_format.space_before = Pt(6)
                ip.paragraph_format.space_after = Pt(4)
                ip.paragraph_format.keep_with_next = True
                
                # Add picture scaled to fit width nicely (max 6.2 inches)
                try:
                    ip.add_run().add_picture(img_path, width=Inches(6.2))
                except Exception as e:
                    print(f"Error adding image {img_path}: {e}")
                    
                if img_caption:
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_after = Pt(14)
                    cap_run = cap_p.add_run(f"Figure: {img_caption}")
                    cap_run.font.name = "Calibri"
                    cap_run.font.size = Pt(9)
                    cap_run.font.italic = True
                    cap_run.font.color.rgb = RGBColor(113, 128, 150)

    doc.save("Review_3_Project_Functionalities.docx")
    print("Saved Review_3_Project_Functionalities.docx successfully!")

if __name__ == "__main__":
    create_docx()
